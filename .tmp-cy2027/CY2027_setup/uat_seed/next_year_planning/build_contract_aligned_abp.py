"""Derive a contract-aligned CY2027 ABP from the supplied FY27 profile.

The source contains the approved day/asset distribution.  This utility scales
only quantities to a documented annual target and adjusts the final row by the
rounding residual, so the target is exact at the ABP's two-decimal MMBtu grain.
"""

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source" / "FY27_ABP_2027_UI_Test.xlsx"
OUTPUT = HERE / "source" / "FY27_ABP_2027_Contract_Aligned_75_00_TBTU.xlsx"
TARGET_MMBTU = Decimal("75000000.00")


def main():
    source = load_workbook(SOURCE, read_only=True, data_only=True)["ABP_Demand"]
    rows = list(source.iter_rows(min_row=2, values_only=True))
    total = sum((Decimal(str(row[2])) for row in rows), Decimal("0"))
    if total <= 0:
        raise ValueError("The supplied ABP has no positive demand.")
    scale = TARGET_MMBTU / total
    output_rows = []
    for day, asset, quantity in rows:
        output_rows.append((day, asset, (Decimal(str(quantity)) * scale).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)))
    residual = TARGET_MMBTU - sum((row[2] for row in output_rows), Decimal("0"))
    day, asset, quantity = output_rows[-1]
    output_rows[-1] = (day, asset, quantity + residual)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ABP_Demand"
    sheet.append(["Date", "AssetCode", "Quantity"])
    for row in output_rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 18
    workbook.save(OUTPUT)
    print(f"Created {OUTPUT.name}: {len(output_rows)} rows, {TARGET_MMBTU} MMBtu / 75.00 TBTU.")


if __name__ == "__main__":
    main()
